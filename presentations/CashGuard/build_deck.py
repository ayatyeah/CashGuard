from pathlib import Path
import math, shutil, json
from PIL import Image
import cairosvg
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT=Path(__file__).resolve().parent
A=ROOT/'assets'; A.mkdir(exist_ok=True)
BRAND=Path('/home/yeahayat/ForteWin-ui/src/frontend/public/brand')
for name in ['cashguard-forte','forte-hero']:
 shutil.copy2(BRAND/f'{name}.svg', A/f'{name}.svg')
 if name != 'cashguard-forte':
  cairosvg.svg2png(url=str(A/f'{name}.svg'),write_to=str(A/f'{name}.png'),output_width=1000)
 # cashguard-forte.png is rendered by Chromium to preserve its luminance mask.
# A subtle radial burgundy treatment derived from the website's hero palette.
w,h=1600,900
im=Image.new('RGB',(w,h)); pix=im.load()
for y in range(h):
 for x in range(w):
  g=math.exp(-(((x-w*.90)/(w*.65))**2+((y-h*.16)/(h*.95))**2)*2)
  pix[x,y]=(int(15+90*g),int(10+0*g),int(18+34*g))
im.save(A/'hero-bg.jpg',quality=95)
P=Presentation();P.slide_width=Inches(16);P.slide_height=Inches(9)
P.core_properties.title='CashGuard — от видео к проверяемому событию'
P.core_properties.subject='ForteWin / кейс ForteBank / Datathon 2026'
P.core_properties.author='ForteWin'
C={'bg':'0B0B0E','panel':'17151B','line':'342831','text':'FAF8FA','muted':'B5A8B4','pink':'EFA7C6','accent':'B91B5C','green':'8BD6AC','amber':'E8BB77'}
notes=[]
def rect(s,x,y,w,h,fill='panel',line=None,radius=True):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
 sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(C.get(fill,fill));sh.line.fill.background()
 if line:sh.line.color.rgb=RGBColor.from_string(C.get(line,line));sh.line.width=Pt(.8)
 if radius:sh.adjustments[0]=.12
 return sh

def txt(s,text,x,y,w,h,size=22,color='text',bold=False):
 box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=box.text_frame
 tf.clear();tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Inter';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(C.get(color,color));p.space_after=Pt(10)
 return box

def pic(s,name,x,y,w=None,h=None):
 return s.shapes.add_picture(str(A/name),Inches(x),Inches(y),width=Inches(w) if w else None,height=Inches(h) if h else None)

def base(kicker,title=None,sub=None,hero=False):
 s=P.slides.add_slide(P.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(C['bg'])
 if hero:pic(s,'hero-bg.jpg',0,0,16,9)
 pic(s,'cashguard-forte.png',.62,.42,.42,.42);txt(s,'CashGuard',1.16,.43,2.2,.4,17,bold=True)
 txt(s,kicker.upper(),4.1,.48,10.9,.35,11,'pink')
 rect(s,.65,8.47,14.7,.012,'line',radius=False)
 txt(s,'FORTEWIN  /  КЕЙС FORTEBANK',.65,8.65,9,.22,9,'muted')
 txt(s,f'{len(P.slides):02d}',14.8,8.61,.6,.28,11,'pink')
 if title:
  title_box=txt(s,title,.7,1.35,14.5,1.3,35 if '\n' in title else 37,bold=True)
  for paragraph in title_box.text_frame.paragraphs: paragraph.space_after=Pt(0)
 if sub:txt(s,sub,.73,2.85 if title and '\n' in title else 2.6,14.3,.4,17,'muted')
 return s

def note(s,text,source):
 t=text+'\n\nИсточники: '+source
 s.notes_slide.notes_text_frame.text=t
 notes.append(f'СЛАЙД {len(P.slides):02d}\n{t}')

def card(s,x,y,w,h,num,title,body,color='pink'):
 rect(s,x,y,w,h,line='line');txt(s,num,x+.28,y+.25,w-.5,.5,18,color,bold=True);txt(s,title,x+.28,y+.95,w-.56,1.0,24,bold=True);txt(s,body,x+.28,y+2,w-.56,h-2.12,18,'muted')

s=base('Datathon 2026 · локальная видеоаналитика',hero=True)
pic(s,'forte-hero.png',11.6,1.4,3.3)
txt(s,'CashGuard',.8,2.05,10,1.35,76,bold=True)
txt(s,'От видео —\nк проверяемому событию',.85,3.7,10.3,1.9,39,bold=True)
txt(s,'Контроль кассовых операций\nс таймкодами, кадрами и нормативным контекстом.',.87,6.12,10.5,1.05,21,'pink')
rect(s,.87,7.5,3.35,.49,'accent');txt(s,'КОМАНДА FORTEWIN',1.1,7.61,3,.3,11,bold=True)
note(s,'15 секунд. CashGuard — наш прототип по кейсу ForteBank. Он помогает проверяющему перейти от длинной записи к конкретным событиям и их доказательствам. Это инструмент поддержки проверки человеком.', 'README.md; Кейс_ForteBank.docx §4–7; дизайн работающего UI ForteWin-ui.')

s=base('01 / задача','На видео есть действие.\nПроверяющему нужен контекст.')
card(s,.75,3.05,4.65,3.8,'01','Найти момент','Просмотреть запись и определить,\nгде начинается значимое событие.')
card(s,5.68,3.05,4.65,3.8,'02','Понять ситуацию','Сопоставить действия до, во время\nи после подозрительного момента.')
card(s,10.61,3.05,4.65,3.8,'03','Обосновать вывод','Собрать кадры, описание\nи релевантную ссылку на правило.')
txt(s,'Ценность CashGuard — собрать этот путь в одном рабочем пространстве.',.8,7.42,14.4,.6,23,'pink')
note(s,'25 секунд. Проблема не только в поиске объекта. Для проверки нужны момент, последовательность и основания. Мы не заявляем измеренного сокращения трудозатрат: это ценность, которую ещё предстоит измерить в пилоте.','Кейс_ForteBank.docx §4–6; UI review workspace.')

s=base('02 / продукт','Интерфейс, который ведёт к проверке')
pic(s,'interface.png',.77,2.6,9.0,5.625)
rect(s,10.12,2.6,5.1,5.62,line='line')
txt(s,'Знакомая визуальная среда',10.47,2.95,4.35,.9,25,bold=True)
txt(s,'Бордовый акцент Forte\nи знак F в плашке.',10.47,4.02,4.3,1,20,'pink')
txt(s,'Загрузка записи\nИстория анализов\nСобытия и доказательства\nЭкспорт результата',10.47,5.27,4.2,2.2,20,'muted')
note(s,'25 секунд. Покажите рабочую среду: загрузка, история и результат в одном месте. Это снимок текущего тёмного интерфейса; для презентации API-ответы обезличены и история пуста. Заявления о производительности из этого экрана не следуют.','ForteWin-ui/src/frontend/src/components/ui/app-1.tsx; index.css; локальный снимок localhost:3000, пустая история.')

s=base('03 / пользовательский путь','Четыре шага от записи к решению','Система подготавливает материал. Окончательную оценку даёт человек.')
steps=[('01','Загрузить','.mov-запись\nкассовой операции'),('02','Проанализировать','Локальные модели\nи проверка кандидатов'),('03','Проверить','Таймкод, кадры,\nописание и норматив'),('04','Сохранить','Результат в истории\nи экспорт JSON')]
for i,(n,t,b) in enumerate(steps):
 x=.75+i*3.72
 rect(s,x,3.62,3.5,3.35,line='line');txt(s,n,x+.25,3.89,3,.6,32,'pink',True);txt(s,t,x+.25,4.84,3,.8,22,bold=True);txt(s,b,x+.25,5.75,3,1.1,18,'muted')
 if i<3:txt(s,'→',x+3.49,4.3,.28,.5,19,'pink')
note(s,'25 секунд. Пользователь загружает видео, сервис создаёт задание, затем показывает результат. Каждый сигнал можно проверить по исходной записи. В текущем продукте есть экспорт JSON; не заявляем автоматизированный окончательный вердикт человека.','src/frontend/src/lib/api.ts; src/backend/internal/transport/http.go; UI app-1.tsx.')

s=base('04 / объяснимый результат','Не только сигнал — основания для проверки','Схема карточки события. Не пример из закрытого датасета.')
rect(s,.75,3.3,8.65,4.5,line='line');rect(s,1.05,3.6,8.05,.64,'382519');txt(s,'ПРИЗНАКИ ВОЗМОЖНОГО НАРУШЕНИЯ',1.3,3.77,7.5,.35,14,'amber',True)
for i,t in enumerate(['ДО СОБЫТИЯ','КЛЮЧЕВОЙ МОМЕНТ','ПОСЛЕ СОБЫТИЯ']):
 x=1.07+i*2.69;rect(s,x,4.56,2.46,1.5,'211B23','line');txt(s,'[ кадр ]',x+.55,4.95,1.4,.45,20,'muted');txt(s,t,x+.13,6.27,2.25,.5,10,'pink')
txt(s,'Тип события  ·  таймкод  ·  наблюдаемый факт',1.12,7.04,7.9,.5,19)
for i,(a,b) in enumerate([('01 / Время','Переход к нужному моменту'),('02 / Доказательства','1–3 кадра для проверки'),('03 / Основание','Описание и ссылка на норматив')]):
 y=3.48+i*1.46;txt(s,a,10, y,5,.45,22,'pink',True);txt(s,b,10,y+.6,5,.65,18,'muted')
note(s,'25 секунд. Карточка связывает сигнал со временем, кадрами и нормативом. На слайде только схема, без реальных кадров и описаний датасета. Норматив в runtime добавляется при наличии пометки source_verified; юридическую корректность трактовки нужно проверять отдельно.','scripts/cashguard_analyzer.py; app/ml/pipeline.py; src/backend/internal/domain/job.go.')

s=base('05 / архитектура','Несколько локальных моделей.\nОдин проверяемый результат.')
rect(s,.75,3.2,14.5,4.75,'111015','line');txt(s,'ЛОКАЛЬНЫЙ КОНТУР ОБРАБОТКИ',1.05,3.45,13,.4,11,'pink',True)
items=[('Веб + API','React / TypeScript\nGo / Gin'),('Восприятие','YOLOX + RTMPose\nТреки и зоны'),('Проверка','CLIP / OWLv2\nQwen3-VL'),('События','Обученный скорер\nДедупликация'),('Результат','JSON + кадры\nНормативы')]
for i,(a,b) in enumerate(items):
 x=1.04+i*2.84;rect(s,x,4.2,2.6,2.17,'panel','line');txt(s,a,x+.18,4.45,2.3,.45,20,bold=True);txt(s,b,x+.18,5.22,2.3,1,16,'muted')
 if i<4:txt(s,'→',x+2.62,4.9,.3,.5,18,'pink')
txt(s,'Хранение заданий: PostgreSQL / SQLite   •   Видео: локальное файловое хранилище',1.1,7.1,13.3,.5,17,'muted')
note(s,'35 секунд. Go запускает локальный Python-анализатор через командный интерфейс. Восприятие выделяет объекты, позы, треки и зоны. Дополнительные модели дают признаки для скорера; набор этапов зависит от настроек и бюджета. VLM не является единственным источником окончательного решения. PostgreSQL используется в Compose, есть адаптер SQLite.','app/ml/{perception,tracking,events,semantic,ovd,vlm,pipeline,scorer}.py; src/backend/internal/infrastructure; compose.yaml.')

s=base('06 / данные и границы','Видео остаётся внутри контура','Локальный инференс после загрузки весов моделей.')
card(s,.75,3.25,4.65,3.85,'ЛОКАЛЬНО','Без облачного AI','Обработка кадров и работа\nмоделей на своём оборудовании.')
card(s,5.68,3.25,4.65,3.85,'НАБЛЮДАЕМО','Без идентификации лиц','Анализ видимых действий,\nобъектов и положения в сцене.')
card(s,10.61,3.25,4.65,3.85,'ПРОВЕРЯЕМО','Человек принимает решение','Сигнал модели — основание\nдля проверки, не санкция.')
txt(s,'Невидимые операции, документы и действия вне кадра нельзя считать установленными фактами.',.8,7.55,14.4,.55,18,'pink')
note(s,'25 секунд. Локальность — требование кейса и свойство ML-архитектуры, а не заявление об аудите безопасности всего сервиса. Отдельно разворачиваются веб-доступ и API. Мы не используем распознавание личности и не делаем выводы об информации, которой нет в кадре.','Кейс_ForteBank.docx §7; offline flags in app/ml/vlm.py and scripts/cashguard_analyzer.py; perception.py.')

s=base('07 / качество','Качество измеряем.\nПропуски не скрываем.','Последний сохранённый отчёт разработки: 40 видео, 17 октября 2026.')
for x,n,l,d in [( .8,'77,5%','Верный вердикт','Групповая проверка разработки'),(5.78,'12 / 21','Найдено позитивных видео','9 пропусков в той же проверке'),(10.76,'0 / 19','Ложных тревог','На негативных видео отчёта')]:
 rect(s,x,3.62,4.45,2.45,line='line');txt(s,n,x+.27,3.9,4,.85,45,'pink',True);txt(s,l,x+.27,4.92,4,.5,18,bold=True);txt(s,d,x+.27,5.57,4,.4,12,'muted')
rect(s,.8,6.44,14.4,1.44,'281322','line');txt(s,'90% на обучающем наборе — не качество на новых видео.',1.08,6.65,13.8,.5,22,bold=True)
txt(s,'Все 40 видео использовались в разработке; независимый тест ещё нужен. Полнота по событиям не измерена.',1.08,7.27,13.8,.45,15,'muted')
note(s,'35 секунд. 77,5% — групповой development CV из сохранённого отчёта, не независимый тест. TP=12, FN=9, TN=19, FP=0; recall позитивных видео 57,1%. Есть риск утечки, поскольку не все обучаемые компоненты переобучались внутри каждого фолда и по этим данным выбирали решения. На обучающем наборе 17/21 позитивных найдено и 19/19 негативных верны — 90%, но это fit-оценка. Приоритет — снизить пропуски и подтвердить результат на новых сценах. В этой работе обучение и оценка повторно не запускались.','artifacts/checks-20261017/comparison.json (candidate.development_cv и training_set); docs/claude_code_next_steps.txt.')

s=base('08 / следующий шаг','Учить последовательность,\nа не запоминать кадр.','План развития — эти этапы ещё предстоит завершить.')
card(s,.75,3.25,4.65,4,'01','Разметить события','Полный локальный просмотр.\nИнтервалы действий, доказательства\nи нормальные альтернативы.')
card(s,5.68,3.25,4.65,4,'02','Добавить контекст','Состояния операции во времени.\nРаздельно: не произошло\nи не было видно.')
card(s,10.61,3.25,4.65,4,'03','Проверить перенос','Новые сцены и ракурсы.\nРазделение по видео и камерам.\nКонтроль пропусков и времени.')
txt(s,'Критерий прогресса: меньше пропусков на новых сценах при контролируемом числе ложных тревог.',.8,7.65,14.4,.5,17,'pink')
note(s,'25 секунд. Улучшение планируем через разметку действий и временной контекст, а не десять одинаковых прогонов. Данные должны оставаться локально. Независимая оценка требует новых постановочных сцен, не тестов жюри. Это roadmap, не уже внедрённая функциональность.','docs/claude_code_next_steps.txt; tests/test_pipeline_budget.py; scripts/train_scorer.py.')

s=base('09 / пилот','Что проверяем перед внедрением','Переход от работающего прототипа к надёжному инструменту.')
rows=[('Качество','Пропуски, ложные тревоги, тип и таймкод каждого события'),('Новые камеры','Устойчивость к ракурсу, освещению и частичным перекрытиям'),('Скорость','Не более 5 минут от загрузки — вместе с ожиданием в очереди'),('Польза','Время ручной проверки и доля подтверждённых сигналов')]
for i,(a,b) in enumerate(rows):
 y=3.32+i*1.05;rect(s,.8,y,14.4,.9,line='line');txt(s,a,1.06,y+.22,3.15,.5,22,'pink',True);txt(s,b,4.38,y+.24,10.5,.5,18)
txt(s,'Условие пилота: согласованные метрики и отдельные данные для оценки.',.85,7.86,14,.4,16,'muted')
note(s,'25 секунд. Полный SLA пока не доказан: в Go worker сейчас таймаут 15 минут, а бюджет Python не гарантирует 5 минут от загрузки при очереди. Поэтому эта строка — критерий пилота, не обещание текущего сервиса. Измеримого ROI пока нет; предлагаем измерять время проверки.','Кейс_ForteBank.docx §10; src/backend/internal/usecase/service.go; scripts/cashguard_analyzer.py; docs/api_contract_gap.md.')

s=base('10 / демонстрация',hero=True)
pic(s,'forte-hero.png',11.65,1.55,3.0)
txt(s,'CashGuard',.85,1.9,10,1.25,66,bold=True)
txt(s,'Каждый сигнал\nможно проверить.',.87,3.45,11,1.9,47,bold=True)
txt(s,'Видео → событие → доказательства → решение человека',.9,6.1,12.7,.8,22,'pink')
rect(s,.9,7.33,5.4,.63,'accent');txt(s,'ДЕМОНСТРАЦИЯ РАБОЧЕГО СЦЕНАРИЯ',1.15,7.52,5,.3,12,bold=True)
note(s,'20 секунд. Завершите основной рассказ и перейдите к локальной демонстрации. Покажите загрузку, готовый результат, переход к таймкоду и норматив. Выберите заранее проверенный сценарий; не утверждайте, что система находит все нарушения. Слайды 12–14 — резерв для вопросов.','Основная часть: слайды 1–11, около 4 минут 35 секунд; оставшееся время — короткая демонстрация.')

s=base('Приложение A / API','Контракт для автоматической проверки','API-ключ передаётся в X-API-Key. Секреты в презентацию не включены.')
rect(s,.8,3.25,8.65,4.4,line='line')
txt(s,'POST  /api/v1/analyze',1.12,3.65,8,.55,24,'pink',True);txt(s,'multipart/form-data: file  →  job_id',1.12,4.35,8,.5,19,'muted')
txt(s,'GET  /api/v1/jobs/{job_id}',1.12,5.27,8,.6,24,'pink',True);txt(s,'status  ·  verdict  ·  violations[]',1.12,6.08,8,.5,19,'muted')
txt(s,'Выход',10.1,3.6,4.8,.5,27,bold=True);txt(s,'Тип и описание\nТаймкод в секундах\n1–3 JPEG-кадра\nНормативная ссылка',10.1,4.48,4.8,2.4,22,'muted')
txt(s,'Бинарный verdict: violation / no_violation. Техническая ошибка должна оставаться failed.',.85,7.97,14.4,.35,15,'pink')
note(s,'Резерв. Маршруты присутствуют в backend. Не путать наличие маршрутов с полностью доказанным SLA: полный прогон очереди и deadline нужен отдельно. Вход .mov; API схемы имеют два вердикта.','src/backend/internal/transport/http.go; src/backend/internal/domain/job.go; Кейс_ForteBank.docx §10.')

s=base('Приложение B / концепт','Три уровня для человека.\nДва вердикта для API.','Предложение для следующей версии интерфейса — пока не внедрено.')
levels=[('01','Признаки не обнаружены','no_violation','green'),('02','Возможное нарушение','violation','amber'),('03','Выраженные признаки','violation','pink')]
for i,(n,t,v,c) in enumerate(levels):
 x=.8+i*4.98;rect(s,x,3.55,4.45,3.4,line='line');txt(s,n,x+.26,3.84,4,.5,25,c,True);txt(s,t,x+.26,4.65,3.95,1.1,26,bold=True);txt(s,'API → '+v,x+.26,6.27,3.95,.4,17,c)
txt(s,'Полнота анализа — отдельный признак. Плохая видимость не доказывает отсутствие нарушения.',.85,7.55,14.2,.65,18,'muted')
note(s,'Резерв. Это UX-концепт для проверки человеком. Не называем третий статус доказанным нарушением. Уверенность должна опираться на наблюдаемые основания и валидацию, а не на произвольные проценты модели. Неполнота не равна нарушению и не должна превращаться в ложную уверенность.','docs/claude_code_next_steps.txt; Кейс_ForteBank.docx §4, §10 и FAQ о недостаточности данных.')

s=base('Приложение C / состояние проекта','Что готово. Что ещё нужно доказать.')
rect(s,.8,3.05,7.05,3.6,line='line');rect(s,8.12,3.05,7.05,3.6,line='line')
txt(s,'Реализовано в коде',1.1,3.36,6.4,.6,27,'pink',True)
txt(s,'Веб-интерфейс и API\nЛокальный ML-пайплайн\nКадры, таймкоды и экспорт\nОбучение, оценка и тесты',1.1,4.2,6.4,2.2,22)
txt(s,'Требует подтверждения',8.42,3.36,6.4,.6,27,'amber',True)
txt(s,'Качество на независимых сценах\nПолнота обнаружения событий\nСтабильные 5 минут с очередью\nЭффект на время работы проверяющего',8.42,4.2,6.4,2.2,21)
txt(s,'Основания презентации',.88,7.02,14,.4,17,bold=True)
txt(s,'ТЗ кейса · код frontend/backend/ML · сохранённый отчёт 17.10.2026 · текущий тёмный UI\nПодробные пути и оговорки — в заметках к слайдам и файле SOURCE_REVIEW.txt.',.88,7.55,14.3,.7,15,'muted')
note(s,'Резерв. Это обзор текущего состояния, не сертификат готовности. Изучены основные подсистемы и структура папки; бинарные веса и закрытые видеозаписи не интерпретировались для презентации. Текущие незакоммиченные изменения пользователя не трогались.','README.md; app/; src/; scripts/; tests/; compose.yaml; configs/; design/; artifacts/checks-20261017/comparison.json.')

P.save(ROOT/'CashGuard_ForteWin.pptx')
(ROOT/'speaker-notes.txt').write_text('\n\n'+'\n\n'.join(notes),encoding='utf-8')
print(f'Saved {len(P.slides)} slides')
